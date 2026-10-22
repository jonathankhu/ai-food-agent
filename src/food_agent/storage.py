"""存照片、寫 Excel、算當日總數。

所有函式都是同步的，而且整支程式只在 asyncio 事件迴圈這一條執行緒裡呼叫它們，
所以兩個動作不會同時改到 Excel，不需要另外加鎖。
"""

import os
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .sanitize import clean_filename_part, write_text_cell

NOT_FOOD_NAME = "不是食物"
FAILED_NAME = "分析失敗"

HEADERS = [
    "日期時間",
    "上傳者",
    "食物名稱",
    "成份",
    "食物重量（公克）",
    "食物熱量（大卡）",
    "當日已攝取熱量（大卡）",
    "照片存放位置",
    "備註",
    "上傳者ID",
    "回覆訊息ID",
]
(
    C_TIME, C_UPLOADER, C_NAME, C_INGREDIENTS, C_WEIGHT, C_KCAL,
    C_DAY_TOTAL, C_PHOTO, C_NOTE, C_UPLOADER_ID, C_REPLY_ID,
) = range(1, len(HEADERS) + 1)

TIME_FORMAT = "%Y-%m-%d %H:%M:%S"
DATE_LENGTH = len("2026-10-20")


@dataclass
class LogRecord:
    when: datetime
    uploader: str
    name: str
    ingredients: str
    weight_g: float | None
    kcal: float | None
    photo_path: str
    note: str
    uploader_id: str


@dataclass
class RowInfo:
    row: int
    uploader_id: str
    name: str
    date: str


def date_of(when: datetime) -> str:
    return when.strftime("%Y-%m-%d")


def _number(value: float | None) -> int | float | None:
    if value is None:
        return None
    value = round(float(value), 1)
    return int(value) if value == int(value) else value


def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


class FoodLog:
    def __init__(self, path: Path):
        self.path = Path(path)

    # ---- 開檔、存檔 ----

    def _load(self):
        if self.path.exists():
            wb = load_workbook(self.path)
            return wb, wb.active
        wb = Workbook()
        ws = wb.active
        ws.title = "紀錄"
        ws.append(HEADERS)
        return wb, ws

    def _save(self, wb) -> None:
        """先存成暫存檔再改名，存到一半斷電也不會弄壞原本的紀錄。"""
        self.path.parent.mkdir(parents=True, exist_ok=True)
        tmp = self.path.with_name(self.path.name + ".tmp")
        wb.save(tmp)
        os.replace(tmp, self.path)

    @staticmethod
    def _row_date(ws, row: int) -> str:
        return str(ws.cell(row, C_TIME).value or "")[:DATE_LENGTH]

    # ---- 當日總數 ----

    @classmethod
    def _recompute_day(cls, ws, date: str) -> float:
        """把這一天每一列的「當日已攝取熱量」重算成累計，回傳當天總數。"""
        running = 0.0
        for row in range(2, ws.max_row + 1):
            if cls._row_date(ws, row) != date:
                continue
            kcal = ws.cell(row, C_KCAL).value
            if _is_number(kcal):
                running += kcal
            ws.cell(row, C_DAY_TOTAL).value = _number(running)
        return running

    def day_total(self, date: str) -> float:
        if not self.path.exists():
            return 0.0
        _wb, ws = self._load()
        total = 0.0
        for row in range(2, ws.max_row + 1):
            if self._row_date(ws, row) == date:
                kcal = ws.cell(row, C_KCAL).value
                if _is_number(kcal):
                    total += kcal
        return total

    def count_for_user(self, uploader_id: str, date: str) -> int:
        """這個人這一天已經有幾列（含失敗與不是食物，因為都用掉了額度）。"""
        if not self.path.exists():
            return 0
        _wb, ws = self._load()
        return sum(
            1
            for row in range(2, ws.max_row + 1)
            if self._row_date(ws, row) == date
            and str(ws.cell(row, C_UPLOADER_ID).value or "") == uploader_id
        )

    # ---- 新增、更正 ----

    def append(self, record: LogRecord) -> tuple[int, float]:
        """在最後面加一列。回傳 (列號, 當天總數)。"""
        wb, ws = self._load()
        row = ws.max_row + 1
        write_text_cell(ws.cell(row, C_TIME), record.when.strftime(TIME_FORMAT))
        write_text_cell(ws.cell(row, C_UPLOADER), record.uploader)
        write_text_cell(ws.cell(row, C_NAME), record.name)
        write_text_cell(ws.cell(row, C_INGREDIENTS), record.ingredients)
        ws.cell(row, C_WEIGHT).value = _number(record.weight_g)
        ws.cell(row, C_KCAL).value = _number(record.kcal)
        write_text_cell(ws.cell(row, C_PHOTO), record.photo_path)
        write_text_cell(ws.cell(row, C_NOTE), record.note)
        write_text_cell(ws.cell(row, C_UPLOADER_ID), record.uploader_id)
        total = self._recompute_day(ws, date_of(record.when))
        self._save(wb)
        return row, total

    def set_reply_id(self, row: int, reply_id: int | str) -> None:
        wb, ws = self._load()
        write_text_cell(ws.cell(row, C_REPLY_ID), str(reply_id))
        self._save(wb)

    def find_by_reply_id(self, reply_id: int | str) -> RowInfo | None:
        if not self.path.exists():
            return None
        _wb, ws = self._load()
        wanted = str(reply_id)
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row, C_REPLY_ID).value or "") == wanted:
                return RowInfo(
                    row=row,
                    uploader_id=str(ws.cell(row, C_UPLOADER_ID).value or ""),
                    name=str(ws.cell(row, C_NAME).value or ""),
                    date=self._row_date(ws, row),
                )
        return None

    def apply_correction(self, row: int, name: str, kcal: float) -> tuple[str, float]:
        """改名稱和熱量，重算那一天。回傳 (日期, 當天總數)。"""
        wb, ws = self._load()
        previous_name = str(ws.cell(row, C_NAME).value or "")
        write_text_cell(ws.cell(row, C_NAME), name)
        ws.cell(row, C_KCAL).value = _number(kcal)
        if previous_name in (NOT_FOOD_NAME, FAILED_NAME):
            ws.cell(row, C_INGREDIENTS).value = None  # 舊成份已經不對了
        date = self._row_date(ws, row)
        total = self._recompute_day(ws, date)
        self._save(wb)
        return date, total


class UnsafePathError(Exception):
    pass


def save_photo(
    photos_dir: Path, data: bytes, ext: str, name_part: str, when: datetime
) -> Path:
    """把照片存進 photos_dir。同一秒同名時依序加 _2、_3……，不會覆蓋。"""
    photos_dir.mkdir(parents=True, exist_ok=True)
    safe_part = clean_filename_part(name_part)
    base = f"{when:%Y%m%d-%H%M%S}_{safe_part}"
    for n in range(1, 1000):
        suffix = "" if n == 1 else f"_{n}"
        path = photos_dir / f"{base}{suffix}.{ext}"
        # >>>>>>>>>> 防護 3 開始：確認算出來的路徑沒有跑出 photos 資料夾 <<<<<<<<<<
        if photos_dir.resolve() not in path.resolve().parents:
            raise UnsafePathError(f"檔案路徑跑出照片資料夾：{path.name}")
        # >>>>>>>>>> 防護 3 結束 <<<<<<<<<<
        try:
            with open(path, "xb") as f:  # 檔案已存在就失敗，不會覆蓋
                f.write(data)
        except FileExistsError:
            continue
        return path
    raise OSError("同一秒內檔名重複太多次")
