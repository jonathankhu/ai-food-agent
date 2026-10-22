from datetime import datetime, timezone
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from food_agent import storage
from food_agent.analyzer import AnalysisError, AnalysisResult
from food_agent.bot import FoodBot, photo_extension
from food_agent.storage import FoodLog

CHANNEL = 111
NOW = datetime(2026, 10, 20, 12, 0, 0, tzinfo=timezone.utc)


class FakeMessage:
    _next_id = 1000

    def __init__(self, author_id=1, content="", attachments=(), reference_id=None,
                 channel=CHANNEL, bot=False, name="小明"):
        FakeMessage._next_id += 1
        self.id = FakeMessage._next_id
        self.author = SimpleNamespace(id=author_id, bot=bot, display_name=name)
        self.channel = SimpleNamespace(id=channel)
        self.content = content
        self.attachments = list(attachments)
        self.reference = SimpleNamespace(message_id=reference_id) if reference_id else None
        self.created_at = NOW
        self.replies = []

    async def reply(self, text, **kwargs):
        sent = FakeMessage(author_id=0, content=text, bot=True)
        self.replies.append((text, kwargs))
        return sent


class FakeAttachment:
    def __init__(self, filename="a.jpg", content_type="image/jpeg", data=b"img"):
        self.filename, self.content_type, self._data = filename, content_type, data

    async def read(self):
        return self._data


def food(**kw):
    return AnalysisResult(**{"is_food": True, "name": "牛肉麵", "ingredients": ["牛肉", "麵"],
                             "weight_g": 550.0, "kcal": 650.0, **kw})


def make_bot(config, outcomes):
    outcomes = list(outcomes)

    async def analyze(photo_bytes, ext, cfg):
        out = outcomes.pop(0)
        if isinstance(out, Exception):
            raise out
        return out

    from food_agent.worker import Worker
    flog = FoodLog(config.log_path)
    return FoodBot(config, flog, Worker(config, flog, analyze))


async def upload(bot, msg):
    await bot.handle_message(msg)
    while not bot.worker.queue.empty():
        job = bot.worker.queue.get_nowait()
        await bot.worker.process(job)
        bot.worker.pending[job.uploader_id] -= 1


def texts(msg):
    return [t for t, _ in msg.replies]


def test_photo_extension():
    assert photo_extension(FakeAttachment("x.JPEG")) == "jpg"
    assert photo_extension(FakeAttachment("x.png", "image/png")) == "png"
    assert photo_extension(FakeAttachment("x.pdf", "application/pdf")) is None
    assert photo_extension(FakeAttachment("x.jpg", "text/plain")) is None
    assert photo_extension(FakeAttachment("x.heic", "image/heic")) is None
    assert photo_extension(FakeAttachment("x.jpg", None)) is None


async def test_food_photo_end_to_end(config):
    bot = make_bot(config, [food()])
    msg = FakeMessage(content="午餐", attachments=[FakeAttachment()])
    await upload(bot, msg)
    assert texts(msg)[0] == "處理中⋯"
    assert "**牛肉麵**" in texts(msg)[1] and "當日已攝取熱量：650 大卡" in texts(msg)[1]
    ws = load_workbook(config.log_path).active
    assert ws.cell(2, storage.C_NOTE).value == "午餐"
    assert ws.cell(2, storage.C_REPLY_ID).value
    assert len(list(config.photos_dir.iterdir())) == 1
    assert all(kw["allowed_mentions"].everyone is False for _, kw in msg.replies)


async def test_non_food_and_failure_are_logged(config):
    bot = make_bot(config, [food(is_food=False, name="貓", weight_g=None, kcal=None),
                            AnalysisError("壞了")])
    cat = FakeMessage(attachments=[FakeAttachment()])
    await upload(bot, cat)
    assert "你傳的不是食物，是**貓**" in texts(cat)[1]
    bad = FakeMessage(attachments=[FakeAttachment()])
    await upload(bot, bad)
    assert "分析失敗" in texts(bad)[1] and "當日已攝取熱量：0 大卡" in texts(bad)[1]
    ws = load_workbook(config.log_path).active
    assert ws.cell(2, storage.C_NAME).value == "不是食物" and ws.cell(2, storage.C_INGREDIENTS).value == "貓"
    assert ws.cell(3, storage.C_NAME).value == "分析失敗" and ws.cell(3, storage.C_KCAL).value is None
    assert len(list(config.photos_dir.iterdir())) == 2


async def test_unexpected_error_does_not_kill_worker(config):
    bot = make_bot(config, [RuntimeError("boom"), food()])
    bad = FakeMessage(attachments=[FakeAttachment()])
    await upload(bot, bad)
    assert "出錯" in texts(bad)[-1]
    ok = FakeMessage(attachments=[FakeAttachment()])
    await upload(bot, ok)
    assert "**牛肉麵**" in texts(ok)[1]


async def test_ignores_other_channels_bots_and_non_photos(config):
    bot = make_bot(config, [])
    for msg in [FakeMessage(channel=999, attachments=[FakeAttachment()]),
                FakeMessage(bot=True, attachments=[FakeAttachment()]),
                FakeMessage(attachments=[FakeAttachment("x.pdf", "application/pdf")]),
                FakeMessage(content="大家午安")]:
        await bot.handle_message(msg)
        assert msg.replies == []
    assert bot.worker.queue.empty()


async def test_multiple_photos_each_get_own_row(config):
    bot = make_bot(config, [food(), food(name="沙拉", kcal=100.0)])
    msg = FakeMessage(attachments=[FakeAttachment(), FakeAttachment()])
    await upload(bot, msg)
    assert texts(msg).count("處理中⋯") == 2
    assert load_workbook(config.log_path).active.max_row == 3


async def test_daily_limit_counts_queued_photos_too(config):  # 上限是 3
    bot = make_bot(config, [food()] * 3)
    msg = FakeMessage(attachments=[FakeAttachment()] * 5)
    await bot.handle_message(msg)  # 5 張一起進來，還沒處理，列也還沒寫
    assert texts(msg).count("處理中⋯") == 3
    assert texts(msg).count("今天的額度用完了，明天再傳。") == 2
    assert bot.worker.queue.qsize() == 3


async def correct(bot, replied, author=1, text="改成：清湯麵 400 大卡"):
    msg = FakeMessage(author_id=author, content=text, reference_id=replied)
    await bot.handle_message(msg)
    return msg


async def test_corrections(config):
    bot = make_bot(config, [food(), food(is_food=False, name="貓", weight_g=None, kcal=None)])
    first = FakeMessage(attachments=[FakeAttachment()])
    await upload(bot, first)
    ws = load_workbook(config.log_path).active
    result_id = int(ws.cell(2, storage.C_REPLY_ID).value)

    assert texts(await correct(bot, result_id, author=2)) == ["只有上傳者能更正這筆。"]
    assert texts(await correct(bot, result_id, text="改成：清湯麵")) == ["請用「改成：名稱 熱量 大卡」的格式。"]
    assert (await correct(bot, 424242)).replies == []  # 回覆的不是 bot 的結果訊息

    ok = await correct(bot, result_id)
    assert texts(ok) == ["已更正：**清湯麵**，400 大卡\n**當日已攝取熱量：400 大卡**"]
    assert load_workbook(config.log_path).active.cell(2, storage.C_NAME).value == "清湯麵"

    cat = FakeMessage(attachments=[FakeAttachment()])
    await upload(bot, cat)
    cat_id = int(load_workbook(config.log_path).active.cell(3, storage.C_REPLY_ID).value)
    fixed = await correct(bot, cat_id, text="改成：沙拉 100 大卡")
    assert "當日已攝取熱量：500 大卡" in texts(fixed)[0]
