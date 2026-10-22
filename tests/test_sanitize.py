from openpyxl import Workbook

from food_agent.sanitize import clean_filename_part, write_text_cell


def test_filename_removes_path_tricks():
    for evil in ["../../x", "a/b", "C:\\x", "....", "", "   ", "a\x00b", "\u202eevil"]:
        out = clean_filename_part(evil)
        assert out
        assert not any(ch in out for ch in '/\\:\x00\u202e')
        assert ".." not in out
        assert not out.startswith(".")


def test_filename_truncates_and_defaults():
    assert len(clean_filename_part("牛" * 100)) == 20
    assert clean_filename_part("") == "未命名"
    assert clean_filename_part("牛肉 麵") == "牛肉_麵"


def test_formula_text_becomes_plain_text():
    ws = Workbook().active
    for i, evil in enumerate(["=1+1", "+1", "-1", "@SUM(1)", "\t=1", "正常"], start=1):
        cell = ws.cell(i, 1)
        write_text_cell(cell, evil)
        assert cell.data_type == "s"
        if evil != "正常":
            assert cell.value == "'" + evil
    assert ws.cell(6, 1).value == "正常"
