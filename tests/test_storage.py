from datetime import datetime

import pytest
from openpyxl import load_workbook

from food_agent import storage
from food_agent.storage import FoodLog, LogRecord, UnsafePathError, save_photo

WHEN = datetime(2026, 10, 20, 12, 34, 56)


def rec(name="牛肉麵", kcal=650.0, uid="1", when=WHEN, **kw):
    base = dict(when=when, uploader="小明", name=name, ingredients="牛肉、麵", weight_g=500,
                kcal=kcal, photo_path="/p.jpg", note="", uploader_id=uid)
    base.update(kw)
    return LogRecord(**base)


def test_appends_instead_of_overwriting_and_accumulates(tmp_path):
    flog = FoodLog(tmp_path / "log.xlsx")
    assert flog.append(rec(kcal=650))[1] == 650
    assert FoodLog(tmp_path / "log.xlsx").append(rec(kcal=300))[1] == 950  # 重新開檔仍往下加
    ws = load_workbook(tmp_path / "log.xlsx").active
    assert ws.max_row == 3
    assert [ws.cell(r, storage.C_DAY_TOTAL).value for r in (2, 3)] == [650, 950]


def test_blank_kcal_rows_do_not_count_and_days_are_separate(tmp_path):
    flog = FoodLog(tmp_path / "log.xlsx")
    flog.append(rec(kcal=650))
    _, total = flog.append(rec(name="不是食物", kcal=None, weight_g=None))
    assert total == 650
    _, next_day = flog.append(rec(kcal=100, when=datetime(2026, 10, 21, 0, 0, 1)))
    assert next_day == 100
    assert flog.day_total("2026-10-20") == 650 and flog.day_total("2026-10-21") == 100


def test_count_for_user_counts_every_row(tmp_path):
    flog = FoodLog(tmp_path / "log.xlsx")
    flog.append(rec(uid="1"))
    flog.append(rec(name="分析失敗", kcal=None, uid="1"))
    flog.append(rec(uid="2"))
    assert flog.count_for_user("1", "2026-10-20") == 2
    assert flog.count_for_user("1", "2026-10-21") == 0


def test_large_ids_survive_as_text(tmp_path):
    flog = FoodLog(tmp_path / "log.xlsx")
    row, _ = flog.append(rec(uid="123456789012345678"))
    flog.set_reply_id(row, 987654321098765432)
    info = flog.find_by_reply_id(987654321098765432)
    assert info.uploader_id == "123456789012345678" and info.row == row
    assert flog.find_by_reply_id(1) is None


def test_correction_recomputes_later_rows(tmp_path):
    flog = FoodLog(tmp_path / "log.xlsx")
    r1, _ = flog.append(rec(kcal=650))
    flog.append(rec(kcal=300))
    date, total = flog.apply_correction(r1, "清湯麵", 400)
    assert (date, total) == ("2026-10-20", 700)
    ws = load_workbook(tmp_path / "log.xlsx").active
    assert [ws.cell(r, storage.C_DAY_TOTAL).value for r in (2, 3)] == [400, 700]
    assert ws.cell(2, storage.C_NAME).value == "清湯麵"
    assert ws.cell(2, storage.C_INGREDIENTS).value == "牛肉、麵"  # 食物列的成份不動


def test_correcting_failed_or_non_food_row_clears_ingredients(tmp_path):
    flog = FoodLog(tmp_path / "log.xlsx")
    row, _ = flog.append(rec(name="不是食物", kcal=None, weight_g=None, ingredients="貓"))
    _, total = flog.apply_correction(row, "牛肉麵", 650)
    assert total == 650
    ws = load_workbook(tmp_path / "log.xlsx").active
    assert ws.cell(row, storage.C_INGREDIENTS).value in (None, "")
    assert ws.cell(row, storage.C_WEIGHT).value is None


def test_formula_text_is_stored_as_text(tmp_path):
    flog = FoodLog(tmp_path / "log.xlsx")
    flog.append(rec(name="=1+1", uploader="@everyone", note="+cmd"))
    ws = load_workbook(tmp_path / "log.xlsx").active
    for col in (storage.C_NAME, storage.C_UPLOADER, storage.C_NOTE):
        assert ws.cell(2, col).data_type == "s"
    assert ws.cell(2, storage.C_NAME).value == "'=1+1"


def test_save_photo_collisions_get_suffix(tmp_path):
    photos = tmp_path / "photos"
    paths = [save_photo(photos, bytes([i]), "jpg", "牛肉麵", WHEN) for i in range(3)]
    assert [p.name for p in paths] == [
        "20261020-123456_牛肉麵.jpg", "20261020-123456_牛肉麵_2.jpg", "20261020-123456_牛肉麵_3.jpg"]
    assert [p.read_bytes() for p in paths] == [b"\x00", b"\x01", b"\x02"]


def test_save_photo_evil_names_stay_inside(tmp_path):
    photos = tmp_path / "photos"
    for evil in ["../../x", "a/b", "C:\\x", "...", "", "/etc/passwd"]:
        path = save_photo(photos, b"x", "jpg", evil, WHEN)
        assert photos.resolve() in path.resolve().parents
    assert {p.name for p in tmp_path.iterdir()} == {"photos"}


def test_containment_check_catches_escape(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "clean_filename_part", lambda text: text)  # 假裝防護 3 清洗壞了
    with pytest.raises(UnsafePathError):
        save_photo(tmp_path / "photos", b"x", "jpg", "a/../../../x", WHEN)
    assert list(tmp_path.glob("**/*.jpg")) == []
